"""
Ground Truth Loader — Load GT requirements and test cases from Excel files

Converts ground truth data into:
1. Structured dicts for evaluation (RequirementEvaluator, TestCaseEvaluator)
2. DSPy training examples for few-shot calibration of the extractors

Supports two GT Excel formats:
- Requirements Excel: columns like Requirement ID, Title, Description, Type
- Test Plan Excel:    columns like Test Case ID, Title, Steps, Expected Result

Usage:
    loader = GroundTruthLoader(gt_dir="groundtruth")
    gt = loader.load_all()
    # gt['requirements'] → List[Dict]
    # gt['test_cases']   → List[Dict]

    # For DSPy few-shot optimization:
    examples = loader.to_dspy_examples()
"""

import json
import re
from pathlib import Path
from typing import Dict, List, Optional, Tuple

try:
    import openpyxl

    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False

try:
    import dspy

    DSPY_AVAILABLE = True
except ImportError:
    DSPY_AVAILABLE = False


# ─────────────────────────────────────────────────────────────────────────────
# Header normalisation helpers
# ─────────────────────────────────────────────────────────────────────────────

def _norm(value) -> str:
    """Normalise a cell value to a lower-case stripped string."""
    return str(value).strip().lower() if value is not None else ""


# Known header aliases for requirements sheets
_REQ_HEADER_MAP = {
    "requirement id": "requirement_id",
    "req id": "requirement_id",
    "req_id": "requirement_id",
    "id": "requirement_id",
    "feature id": "requirement_id",
    "tech id": "requirement_id",
    "title": "title",
    "requirement title": "title",
    "feature": "title",
    "feature name": "title",
    "component": "title",
    "technical requirement": "title",
    "name": "title",
    "description": "description",
    "details": "description",
    "detail": "description",
    "requirements": "description",
    "key responsibilities": "description",
    "type": "type",
    "category": "type",
    "acceptance criteria": "acceptance_criteria",
    "criteria": "acceptance_criteria",
    "acceptance_criteria": "acceptance_criteria",
    "user story": "user_story",
    "user_story": "user_story",
}

# Known header aliases for test-case sheets
_TC_HEADER_MAP = {
    "test case id": "test_case_id",
    "tc id": "test_case_id",
    "test_case_id": "test_case_id",
    "id": "test_case_id",
    "title": "title",
    "test case title": "title",
    "test case name": "title",
    "name": "title",
    "description": "description",
    "objective": "description",
    "test steps": "steps",
    "steps": "steps",
    "test step": "steps",
    "step": "steps",
    "expected result": "expected_result",
    "expected outcome": "expected_result",
    "expected": "expected_result",
    "priority": "priority",
    "type": "type",
    "test type": "type",
    "prerequisite": "prerequisites",
    "prerequisites": "prerequisites",
    "preconditions": "prerequisites",
}

# Headers that indicate this sheet contains test cases (not requirements)
_TC_INDICATORS = {
    "test case id", "tc id", "test_case_id", "test steps", "test step",
    "expected result", "expected outcome",
}

# Headers that indicate this sheet contains requirements
_REQ_INDICATORS = {
    "requirement id", "req id", "req_id", "feature id", "tech id",
    "user story", "acceptance criteria",
}


# ─────────────────────────────────────────────────────────────────────────────
# Main Loader Class
# ─────────────────────────────────────────────────────────────────────────────

class GroundTruthLoader:
    """
    Loads ground truth requirements and test cases from Excel files.

    After calling load_all(), access:
        loader.requirements  → List[Dict] of GT requirements
        loader.test_cases    → List[Dict] of GT test cases
    """

    def __init__(self, gt_dir: str = "groundtruth"):
        self.gt_dir = Path(gt_dir)
        self.requirements: List[Dict] = []
        self.test_cases: List[Dict] = []

    # ── Public API ────────────────────────────────────────────────────────────

    def load_all(self) -> Dict:
        """
        Load all .xlsx files from the GT directory.

        Returns:
            {'requirements': [...], 'test_cases': [...]}
        """
        result: Dict[str, List] = {"requirements": [], "test_cases": []}

        if not OPENPYXL_AVAILABLE:
            print("  ⚠ openpyxl not installed — cannot load Excel GT files")
            print("      Run: pip install openpyxl")
            return result

        if not self.gt_dir.exists():
            print(f"  ⚠ GT directory not found: {self.gt_dir}")
            return result

        xlsx_files = list(self.gt_dir.glob("*.xlsx"))
        if not xlsx_files:
            print(f"  ⚠ No Excel files in: {self.gt_dir}")
            return result

        for xlsx_path in xlsx_files:
            print(f"  Loading GT: {xlsx_path.name}")
            data = self._load_excel(xlsx_path)

            if data["type"] == "requirements":
                result["requirements"].extend(data["items"])
                print(f"    ✓ {len(data['items'])} requirements")
            elif data["type"] == "test_cases":
                result["test_cases"].extend(data["items"])
                print(f"    ✓ {len(data['items'])} test cases")
            else:
                print(f"    ⚠ Could not detect sheet type — skipped")

        self.requirements = result["requirements"]
        self.test_cases = result["test_cases"]

        print(
            f"  GT total: {len(self.requirements)} requirements, "
            f"{len(self.test_cases)} test cases"
        )
        return result

    def to_dspy_examples(self) -> List:
        """
        Convert GT requirements to DSPy training examples.

        These examples calibrate the PrecisionRequirementExtractor so it
        learns the correct granularity and terminology for this project.

        Returns:
            List of dspy.Example objects (empty list if DSPy not available
            or no GT requirements loaded)
        """
        if not DSPY_AVAILABLE:
            print("  ⚠ DSPy not installed — cannot create training examples")
            return []

        if not self.requirements:
            print("  ⚠ No GT requirements loaded. Call load_all() first.")
            return []

        examples = []
        for req in self.requirements:
            title = req.get("title", "").strip()
            description = req.get("description", "").strip()
            if not title:
                continue

            # Build a synthetic document snippet that clearly states the requirement
            synthetic_doc = (
                f"The system shall support {title}. "
                f"{description} "
                f"This is a mandatory requirement for the project."
            ).strip()

            # Expected output from PrecisionRequirementExtractor
            expected_output = json.dumps(
                [
                    {
                        "title": title,
                        "description": description or f"The system shall support {title}.",
                        "type": req.get("type", "Functional"),
                        "source_quote": f"The system shall support {title}.",
                        "extraction_type": "direct_mandate",
                        "confidence": "high",
                        "user_story": req.get(
                            "user_story",
                            f"As a user, I want to {title.lower()} so that I can achieve my goal.",
                        ),
                        "acceptance_criteria": _parse_criteria(
                            req.get("acceptance_criteria", "")
                        ) or [f"Verify {title} works as described"],
                    }
                ],
                ensure_ascii=False,
            )

            example = dspy.Example(
                document_text=synthetic_doc,
                requirements_json=expected_output,
            ).with_inputs("document_text")

            examples.append(example)

        print(f"  ✓ Created {len(examples)} DSPy training examples from GT requirements")
        return examples

    def get_requirement_titles(self) -> List[str]:
        """Return sorted list of GT requirement titles."""
        return sorted(r.get("title", "") for r in self.requirements if r.get("title"))

    def get_test_case_titles(self) -> List[str]:
        """Return sorted list of GT test case titles."""
        return sorted(tc.get("title", "") for tc in self.test_cases if tc.get("title"))

    def summary(self) -> str:
        """Return a human-readable summary of what was loaded."""
        lines = [
            f"Ground Truth Summary ({self.gt_dir})",
            f"  Requirements: {len(self.requirements)}",
            f"  Test Cases:   {len(self.test_cases)}",
        ]
        if self.requirements:
            lines.append("  Requirement titles (first 5):")
            for t in self.get_requirement_titles()[:5]:
                lines.append(f"    • {t}")
        if self.test_cases:
            lines.append("  Test case titles (first 5):")
            for t in self.get_test_case_titles()[:5]:
                lines.append(f"    • {t}")
        return "\n".join(lines)

    # ── Private Helpers ───────────────────────────────────────────────────────

    def _load_excel(self, path: Path) -> Dict:
        """
        Load a single Excel file. Auto-detect whether it contains
        requirements or test cases based on column headers.

        Returns:
            {'type': 'requirements'|'test_cases'|'unknown', 'items': [...]}
        """
        try:
            wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        except Exception as e:
            print(f"    ✗ Cannot open {path.name}: {e}")
            return {"type": "unknown", "items": []}

        all_items = []
        detected_type = "unknown"

        for sheet_name in wb.sheetnames:
            sheet = wb[sheet_name]
            rows = list(sheet.iter_rows(values_only=True))

            if not rows:
                continue

            # Find the actual header row — scan up to the first 3 rows.
            # Some sheets have a section title in row 0 and the real column
            # headers in row 1 (e.g. PTW_Requirements_Document.xlsx).
            header_row_idx = 0
            header_map = []
            sheet_type = "unknown"

            for candidate_row_idx in range(min(3, len(rows))):
                raw_headers = [_norm(h) for h in rows[candidate_row_idx]]
                st, hm = self._detect_sheet_type(raw_headers)
                if st != "unknown":
                    header_row_idx = candidate_row_idx
                    sheet_type = st
                    header_map = hm
                    break

            if sheet_type == "unknown":
                continue

            if detected_type == "unknown":
                detected_type = sheet_type

            # Parse data rows (everything after the header row)
            for row in rows[header_row_idx + 1 :]:
                if not any(cell for cell in row if cell is not None):
                    continue  # Skip entirely empty rows

                row_dict: Dict[str, str] = {}
                for i, canonical_key in enumerate(header_map):
                    if i < len(row):
                        cell_val = row[i]
                        if canonical_key and cell_val is not None:
                            row_dict[canonical_key] = str(cell_val).strip()

                item = self._normalize_row(row_dict, sheet_type)
                if item and item.get("title"):
                    all_items.append(item)

        wb.close()
        return {"type": detected_type, "items": all_items}

    def _detect_sheet_type(
        self, raw_headers: List[str]
    ) -> Tuple[str, List[Optional[str]]]:
        """
        Detect if the sheet is a requirements sheet or test-case sheet.

        Returns:
            (sheet_type, header_map)
            - sheet_type: 'requirements' | 'test_cases' | 'unknown'
            - header_map: list of canonical keys aligned with raw_headers
                          (None where the header is not recognised)
        """
        has_tc_indicator = any(h in _TC_INDICATORS for h in raw_headers)
        has_req_indicator = any(h in _REQ_INDICATORS for h in raw_headers)

        if has_tc_indicator:
            sheet_type = "test_cases"
            alias_map = _TC_HEADER_MAP
        elif has_req_indicator:
            sheet_type = "requirements"
            alias_map = _REQ_HEADER_MAP
        else:
            # Fallback: if there's a "title" column, assume test cases
            if "title" in raw_headers or "name" in raw_headers:
                sheet_type = "test_cases"
                alias_map = _TC_HEADER_MAP
            else:
                return "unknown", []

        header_map = [alias_map.get(h) for h in raw_headers]
        return sheet_type, header_map

    def _normalize_row(self, row_dict: Dict[str, str], sheet_type: str) -> Dict:
        """Normalise a parsed row into a standard requirement or test-case dict."""
        if sheet_type == "requirements":
            return {
                "requirement_id": row_dict.get("requirement_id", ""),
                "title": row_dict.get("title", ""),
                "description": row_dict.get("description", ""),
                "type": row_dict.get("type", "Functional"),
                "acceptance_criteria": row_dict.get("acceptance_criteria", ""),
                "user_story": row_dict.get("user_story", ""),
            }
        else:  # test_cases
            return {
                "test_case_id": row_dict.get("test_case_id", ""),
                "title": row_dict.get("title", ""),
                "description": row_dict.get("description", ""),
                "steps": row_dict.get("steps", ""),
                "expected_result": row_dict.get("expected_result", ""),
                "priority": row_dict.get("priority", ""),
                "type": row_dict.get("type", "Functional"),
                "prerequisites": row_dict.get("prerequisites", ""),
            }


# ─────────────────────────────────────────────────────────────────────────────
# Utility
# ─────────────────────────────────────────────────────────────────────────────

def _parse_criteria(raw: str) -> List[str]:
    """Split acceptance criteria text into a list."""
    if not raw or not raw.strip():
        return []
    # Split on newlines, semicolons, or numbered bullet patterns
    parts = re.split(r"\n|\r\n|;\s*|\d+\.\s+", raw)
    return [p.strip() for p in parts if p.strip()]
