"""
utils.py v4.0 - Shared utilities for BA Pipeline
New in v4.0:
  - DocumentClassifier: auto-classifies Primary BRD vs Supporting docs
  - TokenTracker: tracks input/output token usage across LLM calls
"""
import os
import re
import json
import yaml
import fitz
import logging
import pandas as pd
from pathlib import Path
from docx import Document
from pptx import Presentation
from dataclasses import dataclass, field
from typing import List, Dict, Tuple, Optional, Callable
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.metrics.pairwise import cosine_similarity
from dotenv import load_dotenv
from fnmatch import fnmatch

load_dotenv()
logging.basicConfig(level=logging.INFO, format='%(asctime)s [%(levelname)s] %(message)s')
logger = logging.getLogger(__name__)
ROOT_DIR = Path(__file__).parent.parent

@dataclass
class TestStep:
    step_num: int
    action: str
    expected_result: str = ""
    test_data: str = ""

@dataclass
class BusinessRequirement:
    requirement_id: str
    feature_name: str
    description: str
    system: str = ""
    requirements_text: str = ""
    user_story: str = ""
    acceptance_criteria: List[str] = field(default_factory=list)
    test_steps: List[TestStep] = field(default_factory=list)
    test_scenarios: List[str] = field(default_factory=list)
    category: str = ""
    source_file: str = ""
    page_start: int = 0
    page_end: int = 0
    line_start: int = 0
    line_end: int = 0
    confidence: float = 0.0

@dataclass
class TestCase:
    test_case_id: str
    feature_id: str
    feature_name: str = ""
    description: str = ""
    preconditions: List[str] = field(default_factory=list)
    steps: List[TestStep] = field(default_factory=list)

@dataclass
class DocumentPage:
    page_num: int
    text: str
    line_start: int
    line_end: int

@dataclass
class DocumentInfo:
    filename: str
    full_text: str
    pages: List[DocumentPage] = field(default_factory=list)
    total_pages: int = 0
    total_lines: int = 0
    doc_tier: str = "unknown"  # "primary" or "supporting"

class Config:
    _data = None

    @classmethod
    def load(cls, config_path=None):
        if config_path is None:
            config_path = ROOT_DIR / "config.yaml"
        with open(config_path, 'r') as f:
            cls._data = yaml.safe_load(f)
        return cls._data

    @classmethod
    def get(cls, *keys, default=None):
        if cls._data is None:
            cls.load()
        val = cls._data
        for k in keys:
            if isinstance(val, dict):
                val = val.get(k, default)
            else:
                return default
        return val

# === Token Tracker ===
class TokenTracker:
    """Tracks input/output tokens across LLM calls."""
    def __init__(self):
        self.total_input = 0
        self.total_output = 0
        self.call_count = 0
        self.calls = []

    def track(self, stage, input_tokens=0, output_tokens=0):
        self.total_input += input_tokens
        self.total_output += output_tokens
        self.call_count += 1
        self.calls.append({"stage": stage, "input": input_tokens, "output": output_tokens})

    def estimate_from_text(self, stage, input_text, output_text=""):
        """Estimate tokens from text length (~4 chars per token)."""
        inp = len(input_text) // 4
        out = len(output_text) // 4
        self.track(stage, inp, out)

    def summary(self):
        return {
            "total_input_tokens": self.total_input,
            "total_output_tokens": self.total_output,
            "total_tokens": self.total_input + self.total_output,
            "llm_calls": self.call_count
        }

    def log_summary(self, label=""):
        s = self.summary()
        logger.info(f"TOKEN USAGE {label}: {s['llm_calls']} calls | "
                     f"Input: {s['total_input_tokens']:,} | Output: {s['total_output_tokens']:,} | "
                     f"Total: {s['total_tokens']:,}")

# === Document Classifier ===
class DocumentClassifier:
    """Auto-classifies documents as Primary (BRD) or Supporting (JIRA/Solution/RTM)."""

    # Content heuristics for Primary BRD detection
    PRIMARY_CONTENT_PATTERNS = [
        r'Feature\s*#?\s*\d',           # Feature #1, Feature #2
        r'Feature\s+Name\s*:',          # Feature Name:
        r'Feature\s+ID\s*:',            # Feature ID:
        r'Acceptance\s+Criteria\s*:',   # Acceptance Criteria:
        r'Feature\s+Description\s*:',   # Feature Description:
        r'PTWPH\d+-\d+',               # Feature IDs like PTWPH1-1
    ]

    # Content heuristics for Supporting doc detection
    SUPPORTING_CONTENT_PATTERNS = [
        r'Story\s+Points?\s*:',         # JIRA field
        r'Sprint\s*:',                  # JIRA field
        r'Reporter\s*:',               # JIRA field
        r'Assignee\s*:',               # JIRA field
        r'Priority\s*:\s*(Critical|Major|Minor|Blocker)',
        r'Traceability\s+Matrix',       # RTM
        r'Solution\s+Design',           # Solution doc
        r'Technical\s+Architecture',    # Architecture doc
        r'Sequence\s+Diagram',          # Technical doc
        r'Impacted\s+Systems?\s*:',     # Solution doc
    ]

    @classmethod
    def classify_documents(cls, documents: Dict[str, DocumentInfo]) -> Dict[str, str]:
        """Returns {filename: 'primary' | 'supporting'} for all documents."""
        cfg = Config.load()
        doc_cfg = cfg.get('document_classification', {})
        primary_patterns = doc_cfg.get('primary_patterns', ['*Feature*', '*BRD*'])
        supporting_patterns = doc_cfg.get('supporting_patterns', ['*Solution*', '*Traceability*'])
        jira_pattern = doc_cfg.get('jira_pattern', r'^[A-Z]+-\d+')
        auto_detect = doc_cfg.get('auto_detect', True)

        classification = {}

        for fname, doc_info in documents.items():
            tier = cls._classify_single(
                fname, doc_info, primary_patterns, supporting_patterns,
                jira_pattern, auto_detect
            )
            classification[fname] = tier
            doc_info.doc_tier = tier

        primary = [f for f, t in classification.items() if t == 'primary']
        supporting = [f for f, t in classification.items() if t == 'supporting']
        logger.info(f"Document Classification: {len(primary)} Primary, {len(supporting)} Supporting")
        for f in primary:
            logger.info(f"  [PRIMARY]    {f}")
        for f in supporting:
            logger.info(f"  [SUPPORTING] {f}")

        return classification

    @classmethod
    def _classify_single(cls, filename, doc_info, primary_patterns,
                          supporting_patterns, jira_pattern, auto_detect):
        """Classify a single document."""
        name = filename

        # Step 1: Check explicit config patterns first
        for pattern in primary_patterns:
            if fnmatch(name, pattern):
                return "primary"
        for pattern in supporting_patterns:
            if fnmatch(name, pattern):
                return "supporting"

        # Step 2: JIRA ticket pattern
        if re.match(jira_pattern, name):
            return "supporting"

        # Step 3: Auto-detect from content (if enabled)
        if auto_detect and doc_info.full_text:
            sample = doc_info.full_text[:10000]
            primary_score = sum(1 for p in cls.PRIMARY_CONTENT_PATTERNS if re.search(p, sample, re.IGNORECASE))
            supporting_score = sum(1 for p in cls.SUPPORTING_CONTENT_PATTERNS if re.search(p, sample, re.IGNORECASE))

            if primary_score >= 2 and primary_score > supporting_score:
                return "primary"
            if supporting_score >= 2 and supporting_score > primary_score:
                return "supporting"

        # Step 4: Default — treat as supporting to be safe (no false extractions)
        return "supporting"

# === Validation ===
def validate_gt_file(file_path: Path, expected_sheet: str) -> bool:
    if not file_path.exists():
        raise FileNotFoundError(f"GT file not found: {file_path}")
    xl = pd.ExcelFile(str(file_path))
    if expected_sheet not in xl.sheet_names:
        raise ValueError(f"Sheet '{expected_sheet}' not found in {file_path.name}. Available: {xl.sheet_names}")
    return True

def validate_input_dir(dir_path: Path) -> List[Path]:
    if not dir_path.exists():
        raise FileNotFoundError(f"Input directory not found: {dir_path}")
    supported = {'.pdf', '.docx', '.pptx', '.txt'}
    files = [f for f in dir_path.iterdir() if f.suffix.lower() in supported and not f.name.startswith('~')]
    if not files:
        raise FileNotFoundError(f"No input documents found in {dir_path}")
    logger.info(f"Found {len(files)} input document(s): {[f.name for f in files]}")
    return sorted(files)

def validate_api_key() -> str:
    key = os.getenv('GROQ_API_KEY', '')
    if not key or key == 'gsk_your_api_key_here':
        raise ValueError("GROQ_API_KEY not found or not configured in .env file")
    return key

# === PDF Extractor with Page/Line Tracking ===
class PDFExtractor:

    @staticmethod
    def extract_document(file_path: str) -> DocumentInfo:
        file_path = Path(file_path)
        ext = file_path.suffix.lower()
        if ext == '.pdf':
            return PDFExtractor._extract_pdf(file_path)
        elif ext == '.docx':
            return PDFExtractor._extract_docx(file_path)
        elif ext == '.pptx':
            return PDFExtractor._extract_pptx(file_path)
        elif ext == '.txt':
            return PDFExtractor._extract_txt(file_path)
        return DocumentInfo(filename=file_path.name, full_text="")

    @staticmethod
    def _extract_pdf(file_path: Path) -> DocumentInfo:
        pages = []
        full_lines = []
        global_line = 1
        try:
            doc = fitz.open(str(file_path))
            for pn in range(len(doc)):
                page_text = doc[pn].get_text("text")
                page_lines = page_text.split('\n')
                line_start = global_line
                line_end = global_line + len(page_lines) - 1
                pages.append(DocumentPage(
                    page_num=pn + 1, text=page_text,
                    line_start=line_start, line_end=line_end
                ))
                full_lines.extend(page_lines)
                global_line = line_end + 1
            doc.close()
        except Exception as e:
            logger.error(f"Error reading PDF {file_path}: {e}")

        full_text = ""
        for p in pages:
            full_text += f"\n[PAGE {p.page_num} | LINES {p.line_start}-{p.line_end}]\n{p.text}"

        return DocumentInfo(
            filename=file_path.name, full_text=full_text,
            pages=pages, total_pages=len(pages), total_lines=global_line - 1
        )

    @staticmethod
    def _extract_docx(file_path: Path) -> DocumentInfo:
        try:
            doc = Document(str(file_path))
            lines = [p.text for p in doc.paragraphs if p.text.strip()]
            text = "\n".join(lines)
            page = DocumentPage(page_num=1, text=text, line_start=1, line_end=len(lines))
            return DocumentInfo(
                filename=file_path.name, full_text=f"\n[PAGE 1 | LINES 1-{len(lines)}]\n{text}",
                pages=[page], total_pages=1, total_lines=len(lines)
            )
        except Exception as e:
            logger.error(f"Error reading DOCX {file_path}: {e}")
            return DocumentInfo(filename=file_path.name, full_text="")

    @staticmethod
    def _extract_pptx(file_path: Path) -> DocumentInfo:
        try:
            prs = Presentation(str(file_path))
            pages = []
            global_line = 1
            for i, slide in enumerate(prs.slides, 1):
                slide_lines = []
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text.strip():
                        slide_lines.extend(shape.text.split('\n'))
                slide_text = "\n".join(slide_lines)
                line_start = global_line
                line_end = global_line + len(slide_lines) - 1 if slide_lines else global_line
                pages.append(DocumentPage(
                    page_num=i, text=slide_text,
                    line_start=line_start, line_end=line_end
                ))
                global_line = line_end + 1
            full_text = ""
            for p in pages:
                full_text += f"\n[SLIDE {p.page_num} | LINES {p.line_start}-{p.line_end}]\n{p.text}"
            return DocumentInfo(
                filename=file_path.name, full_text=full_text,
                pages=pages, total_pages=len(pages), total_lines=global_line - 1
            )
        except Exception as e:
            logger.error(f"Error reading PPTX {file_path}: {e}")
            return DocumentInfo(filename=file_path.name, full_text="")

    @staticmethod
    def _extract_txt(file_path: Path) -> DocumentInfo:
        text = file_path.read_text(encoding='utf-8', errors='ignore')
        lines = text.split('\n')
        page = DocumentPage(page_num=1, text=text, line_start=1, line_end=len(lines))
        return DocumentInfo(
            filename=file_path.name,
            full_text=f"\n[PAGE 1 | LINES 1-{len(lines)}]\n{text}",
            pages=[page], total_pages=1, total_lines=len(lines)
        )

    @staticmethod
    def extract_all_documents(dir_path: str) -> Dict[str, DocumentInfo]:
        dir_path = Path(dir_path)
        documents = {}
        supported = {'.pdf', '.docx', '.pptx', '.txt'}
        if not dir_path.exists():
            return documents
        for f in sorted(dir_path.iterdir()):
            if f.suffix.lower() in supported and not f.name.startswith('~'):
                logger.info(f"Extracting: {f.name}")
                doc_info = PDFExtractor.extract_document(str(f))
                if doc_info.full_text.strip():
                    documents[f.name] = doc_info
                    logger.info(f"  -> {doc_info.total_pages} pages, {doc_info.total_lines} lines")
                else:
                    logger.warning(f"  -> Empty document, skipping")
        logger.info(f"Extracted {len(documents)} documents total")
        return documents

    @staticmethod
    def get_combined_text(documents: Dict[str, DocumentInfo]) -> str:
        combined = ""
        for fname, doc in documents.items():
            combined += f"\n\n{'='*60}\n[DOCUMENT: {fname}]\n{'='*60}\n{doc.full_text}"
        return combined

# === Page/Line Locator ===
class SourceLocator:
    @staticmethod
    def locate(feature_name: str, description: str, doc_info: DocumentInfo) -> dict:
        keywords = set()
        for text in [feature_name, description]:
            words = re.findall(r'\b[a-zA-Z]{3,}\b', text.lower())
            keywords.update(words)
        keywords -= {'the', 'and', 'for', 'that', 'this', 'with', 'from', 'have', 'will',
                      'are', 'been', 'should', 'must', 'can', 'into', 'each', 'within',
                      'based', 'using', 'not', 'all', 'also', 'any', 'but', 'has', 'its'}

        if not keywords or not doc_info.pages:
            return {"page_start": 0, "page_end": 0, "line_start": 0, "line_end": 0}

        page_scores = []
        for page in doc_info.pages:
            page_lower = page.text.lower()
            name_match = 1.0 if feature_name.lower() in page_lower else 0.0
            kw_count = sum(1 for kw in keywords if kw in page_lower)
            score = name_match * 100 + kw_count
            page_scores.append((score, page))

        page_scores.sort(key=lambda x: x[0], reverse=True)
        if page_scores[0][0] == 0:
            return {"page_start": 0, "page_end": 0, "line_start": 0, "line_end": 0}

        best_page = page_scores[0][1]
        page_start = best_page.page_num
        page_end = best_page.page_num
        line_start = best_page.line_start
        line_end = best_page.line_end

        threshold = page_scores[0][0] * 0.3
        for score, page in page_scores[1:4]:
            if score >= threshold and abs(page.page_num - page_start) <= 2:
                page_end = max(page_end, page.page_num)
                page_start = min(page_start, page.page_num)
                line_end = max(line_end, page.line_end)
                line_start = min(line_start, page.line_start)

        lines = best_page.text.split('\n')
        fname_lower = feature_name.lower()
        for i, line in enumerate(lines):
            if fname_lower in line.lower() or any(kw in line.lower() for kw in list(keywords)[:3]):
                line_start = best_page.line_start + i
                break

        return {"page_start": page_start, "page_end": page_end,
                "line_start": line_start, "line_end": line_end}

# === GT Readers ===
class BRGroundTruthReader:
    @staticmethod
    def read(file_path: str) -> List[BusinessRequirement]:
        file_path = Path(file_path)
        sheet = Config.get('br_gt', 'sheet_name', default='Requirements Summary')
        header = Config.get('br_gt', 'header_row', default=1)
        cols = Config.get('br_gt', 'columns', default={})
        validate_gt_file(file_path, sheet)
        df = pd.read_excel(str(file_path), sheet_name=sheet, header=header)
        requirements = []
        for _, row in df.iterrows():
            fid = str(row.get(cols.get('feature_id', 'Feature ID'), '')).strip()
            if not fid or fid == 'nan':
                continue
            requirements.append(BusinessRequirement(
                requirement_id=fid,
                feature_name=str(row.get(cols.get('feature_name', 'Feature Name'), '')).strip(),
                description=str(row.get(cols.get('description', 'Description'), '')).strip(),
                system=str(row.get(cols.get('system', 'System'), '')).strip() if pd.notna(row.get(cols.get('system', 'System'))) else '',
                requirements_text=str(row.get(cols.get('requirements', 'Requirements'), '')).strip() if pd.notna(row.get(cols.get('requirements', 'Requirements'))) else '',
            ))
        logger.info(f"Loaded {len(requirements)} BR GT requirements")
        return requirements

class TCGroundTruthReader:
    @staticmethod
    def read(file_path: str) -> Dict[str, TestCase]:
        file_path = Path(file_path)
        if not file_path.exists():
            raise FileNotFoundError(f"TC GT file not found: {file_path}")
        cols = Config.get('tc_gt', 'columns', default={})
        xl = pd.ExcelFile(str(file_path))
        test_cases = {}
        for sheet_name in xl.sheet_names:
            try:
                df = pd.read_excel(xl, sheet_name=sheet_name)
                if df.empty:
                    continue
                col_map = TCGroundTruthReader._detect_columns(df.columns, cols)
                if 'steps' not in col_map:
                    continue
                feature_id, description, preconditions, steps = "", "", [], []
                for _, row in df.iterrows():
                    if not feature_id and 'feature' in col_map and pd.notna(row.get(col_map['feature'])):
                        feature_id = str(row[col_map['feature']]).strip()
                    if not description and 'description' in col_map and pd.notna(row.get(col_map['description'])):
                        description = str(row[col_map['description']]).strip()
                    if 'precondition' in col_map and pd.notna(row.get(col_map['precondition'])):
                        val = str(row[col_map['precondition']]).strip()
                        if val and val.upper() != 'NAN' and val not in preconditions:
                            preconditions.append(val)
                    step_text, result_text = "", ""
                    if 'steps' in col_map and pd.notna(row.get(col_map['steps'])):
                        val = str(row[col_map['steps']]).strip()
                        if val and val.upper() != 'NAN':
                            step_text = val
                    if 'result' in col_map and pd.notna(row.get(col_map['result'])):
                        val = str(row[col_map['result']]).strip()
                        if val and val.upper() != 'NAN':
                            result_text = val
                    if step_text:
                        steps.append(TestStep(step_num=len(steps)+1, action=step_text, expected_result=result_text))
                if feature_id and steps:
                    test_cases[sheet_name] = TestCase(
                        test_case_id=sheet_name, feature_id=feature_id, feature_name=feature_id,
                        description=description, preconditions=preconditions, steps=steps)
            except Exception as e:
                logger.warning(f"Error reading sheet {sheet_name}: {e}")
        logger.info(f"Loaded {len(test_cases)} TC GT test cases")
        return test_cases

    @staticmethod
    def _detect_columns(columns, cfg_cols):
        col_map = {}
        columns_lower = {str(c).lower().strip(): c for c in columns}
        mappings = [
            ('feature', [cfg_cols.get('feature', 'Feature').lower(), 'feature', 'feature id']),
            ('description', [cfg_cols.get('description', 'Description').lower(), 'description']),
            ('steps', [cfg_cols.get('steps', 'Steps').lower(), 'steps', 'step', 'test steps']),
            ('precondition', [cfg_cols.get('precondition', 'Pre-Condition').lower(), 'pre-condition', 'precondition']),
            ('result', [cfg_cols.get('actual_result', 'Actual result').lower(), 'actual result', 'expected result']),
        ]
        for key, candidates in mappings:
            for cand in candidates:
                if cand in columns_lower:
                    col_map[key] = columns_lower[cand]
                    break
            if key not in col_map:
                for col_name, orig in columns_lower.items():
                    for cand in candidates:
                        if cand in col_name:
                            col_map[key] = orig
                            break
                    if key in col_map:
                        break
        return col_map

# === Semantic Matcher ===
class SemanticMatcher:
    def __init__(self):
        self.vectorizer = TfidfVectorizer(stop_words='english', ngram_range=(1, 2), max_features=5000, sublinear_tf=True)

    def compute_similarity(self, t1: str, t2: str) -> float:
        try:
            tfidf = self.vectorizer.fit_transform([t1.lower(), t2.lower()])
            return float(cosine_similarity(tfidf[0:1], tfidf[1:2])[0][0])
        except:
            return 0.0

    def match_lists(self, sources: List[str], targets: List[str], threshold: float = 0.25) -> List[Tuple[int, int, float]]:
        if not sources or not targets:
            return []
        try:
            all_texts = [s.lower() for s in sources] + [t.lower() for t in targets]
            tfidf = self.vectorizer.fit_transform(all_texts)
            sim_matrix = cosine_similarity(tfidf[:len(sources)], tfidf[len(sources):])
            pairs = []
            for i in range(len(sources)):
                for j in range(len(targets)):
                    if sim_matrix[i][j] >= threshold:
                        pairs.append((sim_matrix[i][j], i, j))
            pairs.sort(reverse=True)
            matches, used_src, used_tgt = [], set(), set()
            for score, i, j in pairs:
                if i not in used_src and j not in used_tgt:
                    matches.append((i, j, score))
                    used_src.add(i)
                    used_tgt.add(j)
            return matches
        except Exception as e:
            logger.error(f"Matching error: {e}")
            return []

    def deduplicate(self, texts: List[str], threshold: float = 0.70) -> List[int]:
        if len(texts) <= 1:
            return list(range(len(texts)))
        try:
            tfidf = self.vectorizer.fit_transform([t.lower() for t in texts])
            sim = cosine_similarity(tfidf)
            keep, removed = [], set()
            for i in range(len(texts)):
                if i in removed:
                    continue
                keep.append(i)
                for j in range(i + 1, len(texts)):
                    if j not in removed and sim[i][j] >= threshold:
                        removed.add(j)
            return keep
        except:
            return list(range(len(texts)))

# === Context Extractor ===
class ContextExtractor:
    @staticmethod
    def get_context_from_documents(requirement: BusinessRequirement,
                                    documents: Dict[str, DocumentInfo],
                                    max_chars: int = 6000) -> str:
        keywords = set()
        for text in [requirement.feature_name, requirement.description, requirement.requirements_text]:
            keywords.update(re.findall(r'\b[a-zA-Z]{3,}\b', text.lower()))
        keywords -= {'the', 'and', 'for', 'that', 'this', 'with', 'from', 'have', 'will', 'are',
                      'been', 'should', 'must', 'can', 'into', 'each', 'within', 'based', 'using',
                      'implementation', 'not', 'all', 'also'}

        scored_paras = []
        for fname, doc in documents.items():
            paragraphs = re.split(r'\n\s*\n', doc.full_text)
            for para in paragraphs:
                para = para.strip()
                if len(para) < 20:
                    continue
                para_lower = para.lower()
                name_bonus = 50 if requirement.feature_name.lower() in para_lower else 0
                kw_count = sum(1 for kw in keywords if kw in para_lower)
                score = name_bonus + kw_count
                if score > 0:
                    scored_paras.append((score, para, fname))

        scored_paras.sort(reverse=True, key=lambda x: x[0])
        context = ""
        for score, para, fname in scored_paras:
            if len(context) + len(para) > max_chars:
                break
            context += f"[Source: {fname}]\n{para}\n\n"

        if not context:
            if requirement.source_file in documents:
                context = documents[requirement.source_file].full_text[:max_chars]
            else:
                for doc in documents.values():
                    context += doc.full_text[:max_chars // len(documents)]

        return context.strip()

# === Safe Serialization ===
def safe_str(val) -> str:
    if val is None:
        return ""
    if isinstance(val, str):
        return val
    if isinstance(val, dict):
        return json.dumps(val)
    if isinstance(val, list):
        return "; ".join([safe_str(v) for v in val])
    return str(val)

def safe_str_list(lst) -> List[str]:
    if not lst:
        return []
    result = []
    for item in lst:
        if isinstance(item, str):
            result.append(item)
        elif isinstance(item, dict):
            text = item.get('description', item.get('scenario', item.get('text', '')))
            result.append(str(text) if text else json.dumps(item))
        else:
            result.append(str(item))
    return result

# === Excel Writer ===
class ExcelWriter:
    @staticmethod
    def write_test_cases(test_cases: List[TestCase], output_path: str):
        with pd.ExcelWriter(str(output_path), engine='openpyxl') as writer:
            for tc in test_cases:
                rows = []
                for i, step in enumerate(tc.steps):
                    rows.append({
                        'Feature': tc.feature_id if i == 0 else '',
                        'Description': tc.description if i == 0 else '',
                        'Steps': step.action,
                        'Pre-Condition': ("; ".join(tc.preconditions) if tc.preconditions else step.test_data) if i == 0 else '',
                        'Actual result': step.expected_result
                    })
                if not rows:
                    rows = [{'Feature': tc.feature_id, 'Description': tc.description,
                             'Steps': '', 'Pre-Condition': '', 'Actual result': ''}]
                df = pd.DataFrame(rows)
                safe_name = re.sub(r'[\\/*?\[\]:]', '_', tc.test_case_id)[:31]
                df.to_excel(writer, sheet_name=safe_name, index=False)
        logger.info(f"Written {len(test_cases)} test cases to {output_path}")

    @staticmethod
    def write_requirements(requirements: List[BusinessRequirement], output_path: str):
        rows = []
        for req in requirements:
            steps_text = "; ".join([f"{s.step_num}. {s.action}" for s in req.test_steps]) if req.test_steps else ""
            rows.append({
                'Feature ID': req.requirement_id, 'Feature Name': req.feature_name,
                'Description': req.description, 'System': req.system,
                'Requirements': req.requirements_text, 'User Story': req.user_story,
                'Acceptance Criteria': "; ".join(safe_str_list(req.acceptance_criteria)),
                'Test Steps': steps_text,
                'Test Scenarios': "; ".join(safe_str_list(req.test_scenarios)),
                'Category': req.category, 'Source File': req.source_file,
                'Page Start': req.page_start, 'Page End': req.page_end,
                'Line Start': req.line_start, 'Line End': req.line_end,
                'Confidence': req.confidence
            })
        pd.DataFrame(rows).to_excel(str(output_path), index=False, sheet_name='Extracted Requirements')
        logger.info(f"Written {len(requirements)} requirements to {output_path}")

# === JSON Helpers ===
def safe_json_parse(text: str, fallback=None):
    if fallback is None:
        fallback = []
    text = text.strip()
    try:
        return json.loads(text)
    except:
        pass
    match = re.search(r'```(?:json)?\s*([\s\S]*?)```', text)
    if match:
        try:
            return json.loads(match.group(1).strip())
        except:
            pass
    for pattern in [r'(\[[\s\S]*\])', r'(\{[\s\S]*\})']:
        match = re.search(pattern, text)
        if match:
            try:
                return json.loads(match.group(1))
            except:
                pass
    return fallback

def find_gt_files(directory: str, extension: str = '.xlsx') -> List[Path]:
    dir_path = Path(directory)
    if not dir_path.exists():
        return []
    return sorted([f for f in dir_path.glob(f'*{extension}') if not f.name.startswith('~')])
