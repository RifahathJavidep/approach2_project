import io
import json
import logging
import os
import re
from pathlib import Path
from typing import Dict, Any, List, Optional

import dspy
import fitz
from dotenv import load_dotenv
from PIL import Image
from pptx import Presentation

from .file_router import FileRouter, FileType
from .extractors.docx_extractor import DOCXExtractor
from .extractors.image_extractor import StandaloneImageExtractor
from .extractors.ocr_extractor import LocalOCRExtractor, OCRConfig
from .generators.trained_extractor import (
    BusinessFeatureExtractor,
    FunctionalDetailExtractor,
    WorkflowRequirementExtractor,
    TechnicalRequirementExtractor,
)
from .pipeline import _get_lm

load_dotenv()

logger = logging.getLogger("prism.extraction.manual_dspy")

MIN_CONTENT_CHARS = 30




class ManualDSPyExtractor:
    """
    DSPy-powered single-page requirement extractor.

    Runs the same trained extractors as the main pipeline (Business, Functional,
    Workflow, Technical) against a single document page.
    """

    def __init__(self):
        self.ocr = LocalOCRExtractor(config=OCRConfig())
        self._extractors_initialized = False
        self._business_extractor = None
        self._functional_extractor = None
        self._workflow_extractor = None
        self._technical_extractor = None

    def _init_extractors(self):
        if not self._extractors_initialized:
            self._business_extractor = BusinessFeatureExtractor()
            self._functional_extractor = FunctionalDetailExtractor()
            self._workflow_extractor = WorkflowRequirementExtractor()
            self._technical_extractor = TechnicalRequirementExtractor()
            self._extractors_initialized = True

    def extract_from_page(
        self,
        file_path: str,
        description: str,
        page_no: int,
        layers: Optional[List[str]] = None,
    ) -> Dict[str, Any]:
        path = Path(file_path)
        if not path.exists():
            raise FileNotFoundError(f"File not found: {file_path}")

        text = self._extract_page_text(file_path, page_no)

        if len(text.strip()) < MIN_CONTENT_CHARS:
            return {
                "status": "no_requirements",
                "message": "Page has insufficient text content.",
                "requirements": [],
            }

        enriched_text = f"[USER CONTEXT: {description}]\n\n{text}"

        layers = layers or ["business", "functional", "workflow", "technical"]
        requirements = self._run_dspy_extraction(enriched_text, layers)

        if not requirements:
            return {
                "status": "no_requirements",
                "message": "No matching requirements found on this page.",
                "requirements": [],
            }

        relevant = self._filter_by_description(requirements, description)

        if not relevant:
            return {
                "status": "no_requirements",
                "message": f"Extracted {len(requirements)} candidates but none matched '{description}'.",
                "requirements": [],
            }

        return {
            "status": "success",
            "requirements": relevant,
            "total_extracted": len(requirements),
            "total_relevant": len(relevant),
            "source_page": page_no,
            "source_file": path.name,
            "extraction_method": "dspy_trained",
        }

    def _extract_page_text(self, file_path: str, page_no: int) -> str:
        file_type = FileRouter.classify(file_path)

        if file_type in [FileType.PDF, FileType.SCANNED_PDF]:
            doc = fitz.open(file_path)
            if page_no < 1 or page_no > len(doc):
                doc.close()
                raise ValueError(f"Page {page_no} out of range (1-{len(doc)})")

            page = doc[page_no - 1]
            text = page.get_text().strip()

            if len(text) < 100:
                pix = page.get_pixmap(dpi=200)
                img = Image.open(io.BytesIO(pix.tobytes("png")))
                ocr_result = self.ocr.extract_text_from_pil_image(img)
                text = ocr_result.get("text", "")
            doc.close()
            return text

        elif file_type == FileType.PPTX:
            prs = Presentation(file_path)
            if page_no < 1 or page_no > len(prs.slides):
                raise ValueError(f"Slide {page_no} out of range (1-{len(prs.slides)})")
            slide = prs.slides[page_no - 1]
            return "\n".join([s.text for s in slide.shapes if hasattr(s, "text")])

        elif file_type == FileType.DOCX:
            extractor = DOCXExtractor()
            return extractor.extract_text(file_path)

        elif file_type == FileType.IMAGE:
            extractor = StandaloneImageExtractor(ocr_extractor=self.ocr)
            return extractor.extract_text(file_path)

        return ""

    def _run_dspy_extraction(self, text: str, layers: List[str]) -> List[Dict]:
        self._init_extractors()

        layer_map = {
            "business": self._business_extractor,
            "functional": self._functional_extractor,
            "workflow": self._workflow_extractor,
            "technical": self._technical_extractor,
        }

        all_requirements = []

        lm = _get_lm()
        with dspy.context(lm=lm):
            for layer_name in layers:
                extractor = layer_map.get(layer_name)
                if not extractor:
                    continue

                try:
                    result = extractor.forward(document_text=text)
                    reqs = result.get("requirements", [])
                    for req in reqs:
                        req["_extraction_layer"] = layer_name
                    all_requirements.extend(reqs)
                except Exception as e:
                    logger.warning("Layer '%s' failed: %s", layer_name, e, exc_info=True)

        return all_requirements

    def _filter_by_description(
        self, requirements: List[Dict], description: str, threshold: float = 0.25
    ) -> List[Dict]:
        """Filter requirements by keyword overlap with the user's description."""
        desc_words = set(
            w.lower()
            for w in description.split()
            if len(w) > 3 and w.lower() not in {
                "the", "and", "for", "with", "from", "that", "this",
                "must", "should", "system", "user", "want", "able",
            }
        )

        if not desc_words:
            return requirements

        scored = []
        for req in requirements:
            title = req.get("title", "").lower()
            desc = req.get("description", "").lower()
            combined = f"{title} {desc}"

            matched = sum(1 for w in desc_words if w in combined)
            score = matched / len(desc_words) if desc_words else 0

            if score >= threshold:
                req["_relevance_score"] = round(score, 2)
                scored.append(req)

        scored.sort(key=lambda r: r.get("_relevance_score", 0), reverse=True)
        return scored
