"""
Separate Business Requirements (BRs) and Test Cases from 6 project documents.
==============================================================================
Output:
  output/new_project/separated_brs.json       — All extracted BRs
  output/new_project/separated_testcases.json  — All extracted test cases/scenarios
  output/new_project/summary.json              — Counts & source mapping
"""

import json
import subprocess
import re
import os
import openpyxl
from pathlib import Path
from docx import Document
from pptx import Presentation

BASE_DIR = Path(__file__).parent
DOC_DIR = BASE_DIR / "OneDrive_1_3-5-2026 3"
OUTPUT_DIR = BASE_DIR / "output" / "new_project"
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)


# ──────────────────────────────────────────────────────────────
# HELPERS
# ──────────────────────────────────────────────────────────────

def extract_doc_text(filepath):
    """Extract text from legacy .doc using macOS textutil."""
    result = subprocess.run(
        ['textutil', '-convert', 'txt', '-stdout', str(filepath)],
        capture_output=True, text=True
    )
    return result.stdout


# ──────────────────────────────────────────────────────────────
# JIRA .DOC PARSER  (Committed + Uncommitted Feature Lists)
# ──────────────────────────────────────────────────────────────

def parse_jira_doc(filepath, label):
    """Parse JIRA feature export .doc — extract BRs & test scenarios."""
    text = extract_doc_text(filepath)
    filename = os.path.basename(filepath)
    
    blocks = re.split(r'(\[BBMSA-\d+\])', text)
    
    brs = []
    test_cases = []
    
    for i in range(1, len(blocks), 2):
        feature_id = blocks[i].strip('[]')
        if i + 1 >= len(blocks):
            continue
        content = blocks[i + 1]
        
        # ── Extract Feature Title ──
        title_match = re.search(r'^(.*?)(?:Created:)', content, re.DOTALL)
        title = title_match.group(1).strip()[:200] if title_match else "Unknown"
        title = re.sub(r'\s+', ' ', title)
        
        # ── Extract Benefit Hypothesis ──
        bh_text = ""
        bh_match = re.search(
            r'Benefit Hypothesis:\s*(.*?)(?=Acceptance Criteria|How Will We Test|Feature Description|Uncommitted Objective|Exclude From Release|\[BBMSA-|\Z)',
            content, re.DOTALL
        )
        if bh_match:
            bh_text = bh_match.group(1).strip()
            if bh_text and 'None entered' in bh_text[:50]:
                bh_text = ""
        
        # ── Extract Acceptance Criteria ──
        ac_text = ""
        ac_match = re.search(
            r'Acceptance Criteria:\s*(.*?)(?=How Will We Test|Feature Description|Benefit Hypothesis|Uncommitted Objective|Exclude From Release|\[BBMSA-|\Z)',
            content, re.DOTALL
        )
        if ac_match:
            ac_text = ac_match.group(1).strip()
            if 'None entered' in ac_text[:50]:
                ac_text = ""
        
        # ── Extract Feature Description ──
        fd_text = ""
        fd_match = re.search(
            r'Feature Description:\s*(.*?)(?=Acceptance Criteria|How Will We Test|Benefit Hypothesis|Uncommitted Objective|Exclude From Release|\[BBMSA-|\Z)',
            content, re.DOTALL
        )
        if fd_match:
            fd_text = fd_match.group(1).strip()
        
        # ── Extract "How Will We Test" ──
        test_text = ""
        test_match = re.search(
            r'How Will We Test the Feature\?\s*(.*?)(?=Uncommitted Objective|Exclude From Release|Do We Have a Pre-Prod|\Z)',
            content, re.DOTALL
        )
        if test_match:
            test_text = test_match.group(1).strip()
            # Remove template placeholder text
            if test_text.startswith("Please specify testing required"):
                test_text = ""
        
        # ── Build BR entry ──
        br_description = ""
        if fd_text:
            br_description = fd_text
        elif bh_text:
            br_description = bh_text
        
        if br_description or ac_text:
            br = {
                "requirement_id": feature_id,
                "title": title,
                "description": br_description,
                "acceptance_criteria": ac_text,
                "type": "Business Feature",
                "source_file": filename,
                "source_label": label,
            }
            brs.append(br)
        
        # ── Build Test Case entry (only if real content) ──
        if test_text and len(test_text) > 30:
            # Clean up — check if it has actual test steps/scenarios
            has_steps = any(kw in test_text.lower() for kw in [
                'scenario', 'verify', 'validate', 'log in', 'click',
                'create', 'check', 'unit test', 'integration test',
                'manual test', 'automated test', 'e2e', 'uat', 'dvt',
                'given', 'when', 'then', 'expected'
            ])
            
            tc = {
                "test_id": f"TC-{feature_id}",
                "related_requirement": feature_id,
                "title": f"Test Scenarios for: {title}",
                "test_content": test_text,
                "has_detailed_steps": has_steps and len(test_text) > 100,
                "source_file": filename,
                "source_label": label,
            }
            test_cases.append(tc)
        
        # ── Also extract Given/When/Then from Acceptance Criteria ──
        if ac_text and re.search(r'(?:Given|When|Then)', ac_text, re.IGNORECASE):
            gwt_scenarios = re.findall(
                r'(\d+\s+.*?(?:Given|When|Then).*?)(?=\d+\s+|\Z)',
                ac_text, re.DOTALL | re.IGNORECASE
            )
            if gwt_scenarios:
                tc_gwt = {
                    "test_id": f"TC-AC-{feature_id}",
                    "related_requirement": feature_id,
                    "title": f"Given/When/Then Scenarios from AC: {title}",
                    "test_content": ac_text,
                    "has_detailed_steps": True,
                    "source_file": filename,
                    "source_label": label,
                    "note": "Extracted from Acceptance Criteria (Given/When/Then format)"
                }
                test_cases.append(tc_gwt)
    
    return brs, test_cases


# ──────────────────────────────────────────────────────────────
# EXCEL PARSERS
# ──────────────────────────────────────────────────────────────

def parse_happy_town_xlsx(filepath):
    """Parse Happy Town Reports Requirements."""
    wb = openpyxl.load_workbook(str(filepath))
    filename = os.path.basename(filepath)
    brs = []
    
    # Sheet 1: Report Req layout
    ws = wb['Report Req layout']
    current_report = None
    report_num = 0
    sub_items = []
    
    for row in ws.iter_rows(min_row=2, values_only=True):
        report_desc = str(row[4]).strip() if row[4] else ''
        req_text = str(row[5]).strip() if row[5] else ''
        timing = str(row[8]).strip() if len(row) > 8 and row[8] else ''
        security = str(row[10]).strip() if len(row) > 10 and row[10] else ''
        
        if report_desc and report_desc != 'None':
            # Save previous report
            if current_report:
                brs.append({
                    "requirement_id": f"HT-RPT-{report_num:03d}",
                    "title": current_report['title'],
                    "description": current_report['req_text'],
                    "acceptance_criteria": "\n".join(sub_items) if sub_items else "",
                    "type": "Report Requirement",
                    "timing": current_report.get('timing', ''),
                    "security": current_report.get('security', ''),
                    "source_file": filename,
                    "source_sheet": "Report Req layout",
                })
            
            report_num += 1
            current_report = {
                "title": report_desc,
                "req_text": req_text if req_text != 'None' else '',
                "timing": timing if timing != 'None' else '',
                "security": security if security != 'None' else '',
            }
            sub_items = []
        elif req_text and req_text != 'None':
            sub_items.append(req_text)
    
    # Save last report
    if current_report:
        brs.append({
            "requirement_id": f"HT-RPT-{report_num:03d}",
            "title": current_report['title'],
            "description": current_report['req_text'],
            "acceptance_criteria": "\n".join(sub_items) if sub_items else "",
            "type": "Report Requirement",
            "timing": current_report.get('timing', ''),
            "security": current_report.get('security', ''),
            "source_file": filename,
            "source_sheet": "Report Req layout",
        })
    
    # Sheet 2: Secure Rpt Req
    ws2 = wb['Secure Rpt Req ']
    secure_text = []
    for row in ws2.iter_rows(values_only=True):
        if row[0] and str(row[0]).strip():
            secure_text.append(str(row[0]).strip())
    
    if secure_text:
        brs.append({
            "requirement_id": "HT-SEC-001",
            "title": "Happy Town Secure Report Requirements",
            "description": "\n".join(secure_text),
            "acceptance_criteria": "",
            "type": "Security Requirement",
            "source_file": filename,
            "source_sheet": "Secure Rpt Req",
        })
    
    return brs


def parse_incident_xlsx(filepath):
    """Parse Incident Tickets ORD Town."""
    wb = openpyxl.load_workbook(str(filepath))
    filename = os.path.basename(filepath)
    brs = []
    
    # Sheet: Notes — WLAN/DLAN requirements
    ws_notes = wb['Notes']
    notes_text = []
    for row in ws_notes.iter_rows(values_only=True):
        if row[0] and str(row[0]).strip():
            notes_text.append(str(row[0]).strip())
    
    if notes_text:
        brs.append({
            "requirement_id": "INC-NOTES-001",
            "title": "Incident Ticketing System Requirements (WLAN/DLAN)",
            "description": "\n".join(notes_text),
            "acceptance_criteria": "",
            "type": "System Requirement",
            "source_file": filename,
            "source_sheet": "Notes",
        })
    
    # Sheet: Incident Record — field specs
    ws_ir = wb['Incident Record']
    fields = []
    for row in ws_ir.iter_rows(min_row=2, values_only=True):
        en = str(row[0]).strip() if row[0] else ''
        fr = str(row[1]).strip() if row[1] else ''
        notes = str(row[2]).strip() if row[2] else ''
        mapping = str(row[3]).strip() if len(row) > 3 and row[3] else ''
        
        if en and en != 'None':
            fields.append(f"- {en} ({fr}): {notes}" if notes and notes != 'None' else f"- {en} ({fr})")
    
    if fields:
        brs.append({
            "requirement_id": "INC-FIELDS-001",
            "title": "Incident Record Field Specifications",
            "description": "The Incident Record form must include the following fields:\n" + "\n".join(fields),
            "acceptance_criteria": "",
            "type": "UI/System Design Requirement",
            "source_file": filename,
            "source_sheet": "Incident Record",
        })
    
    # Sheet: Search Form
    ws_sf = wb['Search Form']
    search_fields = []
    for row in ws_sf.iter_rows(min_row=2, values_only=True):
        en = str(row[0]).strip() if row[0] else ''
        searchable = str(row[3]).strip() if len(row) > 3 and row[3] else ''
        result = str(row[4]).strip() if len(row) > 4 and row[4] else ''
        
        if en and en != 'None' and searchable:
            search_fields.append(f"- {en}: searchable={searchable}, in_results={result}")
    
    if search_fields:
        brs.append({
            "requirement_id": "INC-SEARCH-001",
            "title": "Incident Search Form Requirements",
            "description": "The search form must support the following fields:\n" + "\n".join(search_fields),
            "acceptance_criteria": "",
            "type": "UI/System Design Requirement",
            "source_file": filename,
            "source_sheet": "Search Form",
        })
    
    # Sheet: Workflow
    ws_wf = wb['Workflow']
    wf_items = []
    for row in ws_wf.iter_rows(values_only=True):
        if row[0] and str(row[0]).strip():
            wf_items.append(str(row[0]).strip())
    
    if wf_items:
        brs.append({
            "requirement_id": "INC-WF-001",
            "title": "Incident Workflow Requirements",
            "description": "\n".join(wf_items),
            "acceptance_criteria": "",
            "type": "Workflow Requirement",
            "source_file": filename,
            "source_sheet": "Workflow",
        })
    
    # Sheet: Notifications
    ws_notif = wb['Notifications']
    notif_items = []
    for row in ws_notif.iter_rows(values_only=True):
        if row[0] and str(row[0]).strip():
            notif_items.append(str(row[0]).strip())
    
    if notif_items:
        brs.append({
            "requirement_id": "INC-NOTIF-001",
            "title": "Incident Notification Requirements",
            "description": "\n".join(notif_items),
            "acceptance_criteria": "",
            "type": "Notification Requirement",
            "source_file": filename,
            "source_sheet": "Notifications",
        })
    
    # Sheet: Search Field Types
    ws_sft = wb['Search Field types (general)']
    sft_items = []
    for row in ws_sft.iter_rows(min_row=3, values_only=True):
        ft = str(row[0]).strip() if row[0] else ''
        widget = str(row[1]).strip() if len(row) > 1 and row[1] else ''
        ops = str(row[2]).strip() if len(row) > 2 and row[2] else ''
        
        if ft and ft != 'None':
            sft_items.append(f"- {ft}: widget={widget}, operations={ops}")
    
    if sft_items:
        brs.append({
            "requirement_id": "INC-SFT-001",
            "title": "Search Field Type Definitions",
            "description": "The search form must support the following field types:\n" + "\n".join(sft_items),
            "acceptance_criteria": "",
            "type": "Technical Requirement",
            "source_file": filename,
            "source_sheet": "Search Field types",
        })
    
    return brs


# ──────────────────────────────────────────────────────────────
# PPTX PARSER
# ──────────────────────────────────────────────────────────────

def parse_pptx(filepath):
    """Parse POS Upgrade Requirement Sample."""
    prs = Presentation(str(filepath))
    filename = os.path.basename(filepath)
    brs = []
    
    # Slide 2 — project highlights
    slide2 = prs.slides[1]
    highlights = []
    for shape in slide2.shapes:
        if hasattr(shape, 'text') and shape.text.strip():
            highlights.append(shape.text.strip())
    
    if highlights:
        brs.append({
            "requirement_id": "POS-001",
            "title": "POS Upgrade Project Highlights",
            "description": "\n".join(highlights),
            "acceptance_criteria": "",
            "type": "Business Feature",
            "source_file": filename,
            "source_slide": 2,
        })
    
    # Slides 3-4 — Process flows (current and proposed)
    for slide_idx, label in [(2, "Current Flow"), (3, "Proposed Flow")]:
        slide = prs.slides[slide_idx]
        programs = []
        for shape in slide.shapes:
            if hasattr(shape, 'text') and shape.text.strip():
                programs.append(shape.text.strip())
        
        if programs:
            brs.append({
                "requirement_id": f"POS-FLOW-{slide_idx+1:03d}",
                "title": f"POS {label} — Programs and Tables",
                "description": "\n".join(programs),
                "acceptance_criteria": "",
                "type": "Technical Requirement",
                "source_file": filename,
                "source_slide": slide_idx + 1,
            })
    
    # Slides 5-8 — Report modifications
    report_titles = {
        4: "Modified Report — Access/Visa Transactions",
        5: "Modified Report — Switch Transactions",
        6: "New Error Report — Swipe Indication Translation Failures",
        7: "Settlement File Specification",
    }
    for slide_idx, rtitle in report_titles.items():
        slide = prs.slides[slide_idx]
        texts = []
        for shape in slide.shapes:
            if hasattr(shape, 'text') and shape.text.strip():
                texts.append(shape.text.strip())
        
        brs.append({
            "requirement_id": f"POS-RPT-{slide_idx+1:03d}",
            "title": rtitle,
            "description": "\n".join(texts) if texts else "(Diagram/screenshot — see slide for visual)",
            "acceptance_criteria": "",
            "type": "Report Requirement",
            "source_file": filename,
            "source_slide": slide_idx + 1,
        })
    
    # Slide 9 — New/Modified Objects table
    slide9 = prs.slides[8]
    objects_list = []
    for shape in slide9.shapes:
        if shape.has_table:
            table = shape.table
            for row_idx, row in enumerate(table.rows):
                if row_idx == 0:
                    continue  # skip header
                cells = [cell.text.strip() for cell in row.cells]
                if cells[1]:  # Object Name exists
                    objects_list.append(
                        f"- {cells[0]}. {cells[1]}: {cells[2]} (Type: {cells[3]}, {cells[4]})"
                    )
    
    if objects_list:
        brs.append({
            "requirement_id": "POS-OBJ-001",
            "title": "POS Upgrade — New/Modified Objects",
            "description": "The following objects must be created or modified:\n" + "\n".join(objects_list),
            "acceptance_criteria": "",
            "type": "Technical Requirement",
            "source_file": filename,
            "source_slide": 9,
        })
    
    return brs


# ──────────────────────────────────────────────────────────────
# DOCX PARSER
# ──────────────────────────────────────────────────────────────

def parse_docx(filepath):
    """Parse Project Objective Sample 3 Snippet."""
    doc = Document(str(filepath))
    filename = os.path.basename(filepath)
    brs = []
    
    # Get objective
    objective = ""
    for para in doc.paragraphs:
        if para.text.strip() and para.style.name != 'Heading 1':
            objective += para.text.strip() + "\n"
    
    if objective:
        brs.append({
            "requirement_id": "PROJ-OBJ-001",
            "title": "Project Objective",
            "description": objective.strip(),
            "acceptance_criteria": "",
            "type": "Business Feature",
            "source_file": filename,
        })
    
    # Tables
    for t_idx, table in enumerate(doc.tables):
        header = [cell.text.strip() for cell in table.rows[0].cells]
        table_title = header[1] if len(header) > 1 else f"Table {t_idx+1}"
        
        items = []
        for row in table.rows[1:]:
            cells = [cell.text.strip() for cell in row.cells]
            if cells[0]:
                items.append(f"- [{cells[0]}] {cells[1]}")
        
        if items:
            scope_type = "In-Scope" if t_idx == 0 else "Out-of-Scope"
            brs.append({
                "requirement_id": f"PROJ-SCOPE-{t_idx+1:03d}",
                "title": f"Project Scope — {scope_type} Items",
                "description": f"{table_title}:\n" + "\n".join(items),
                "acceptance_criteria": "",
                "type": "Scope Definition",
                "source_file": filename,
            })
    
    return brs


# ──────────────────────────────────────────────────────────────
# MAIN
# ──────────────────────────────────────────────────────────────

def main():
    print("=" * 70)
    print("🔀 SEPARATING BRs AND TEST CASES FROM ALL 6 DOCUMENTS")
    print("=" * 70)
    
    all_brs = []
    all_test_cases = []
    
    # ── 1. Committed Feature List (.doc) ──
    print("\n[1/6] Committed Feature List (.doc)...")
    brs1, tc1 = parse_jira_doc(
        DOC_DIR / "2026+PI1+Committed+Feature+List+(DOQ+JIRA)+2026-02-24T13_07_28-0500.doc",
        "Committed"
    )
    all_brs.extend(brs1)
    all_test_cases.extend(tc1)
    print(f"  → {len(brs1)} BRs, {len(tc1)} test scenarios")
    
    # ── 2. Uncommitted Feature List (.doc) ──
    print("[2/6] Uncommitted Feature List (.doc)...")
    brs2, tc2 = parse_jira_doc(
        DOC_DIR / "2026+PI1+Uncommitted+Feature+List+(DOQ+JIRA)+2026-02-24T14_24_46-0500.doc",
        "Uncommitted"
    )
    all_brs.extend(brs2)
    all_test_cases.extend(tc2)
    print(f"  → {len(brs2)} BRs, {len(tc2)} test scenarios")
    
    # ── 3. Happy Town Reports (.xlsx) ──
    print("[3/6] Happy Town Reports Requirements (.xlsx)...")
    brs3 = parse_happy_town_xlsx(DOC_DIR / "hAPPY tOWN rEPORTS REQUIREMENTS.xlsx")
    all_brs.extend(brs3)
    print(f"  → {len(brs3)} BRs, 0 test cases")
    
    # ── 4. Incident Tickets (.xlsx) ──
    print("[4/6] Incident Tickets ORD Town (.xlsx)...")
    brs4 = parse_incident_xlsx(DOC_DIR / "Incident Tickets ORD Town .xlsx")
    all_brs.extend(brs4)
    print(f"  → {len(brs4)} BRs, 0 test cases")
    
    # ── 5. POS Upgrade (.pptx) ──
    print("[5/6] POS Upgrade Requirement Sample (.pptx)...")
    brs5 = parse_pptx(DOC_DIR / "POS upgrade Requirement sample.pptx")
    all_brs.extend(brs5)
    print(f"  → {len(brs5)} BRs, 0 test cases")
    
    # ── 6. Project Objective (.docx) ──
    print("[6/6] Project Objective Sample 3 Snippet (.docx)...")
    brs6 = parse_docx(DOC_DIR / "Project Objective Sample 3 snippet .docx")
    all_brs.extend(brs6)
    print(f"  → {len(brs6)} BRs, 0 test cases")
    
    # ── Save BRs ──
    br_path = OUTPUT_DIR / "separated_brs.json"
    with open(br_path, "w") as f:
        json.dump(all_brs, f, indent=2, ensure_ascii=False)
    
    # ── Save Test Cases ──
    tc_path = OUTPUT_DIR / "separated_testcases.json"
    with open(tc_path, "w") as f:
        json.dump(all_test_cases, f, indent=2, ensure_ascii=False)
    
    # ── Save Summary ──
    summary = {
        "total_brs": len(all_brs),
        "total_test_cases": len(all_test_cases),
        "test_cases_with_detailed_steps": sum(1 for tc in all_test_cases if tc.get("has_detailed_steps")),
        "by_source": {
            "Committed Feature List (.doc)": {"brs": len(brs1), "test_cases": len(tc1)},
            "Uncommitted Feature List (.doc)": {"brs": len(brs2), "test_cases": len(tc2)},
            "Happy Town Reports (.xlsx)": {"brs": len(brs3), "test_cases": 0},
            "Incident Tickets (.xlsx)": {"brs": len(brs4), "test_cases": 0},
            "POS Upgrade (.pptx)": {"brs": len(brs5), "test_cases": 0},
            "Project Objective (.docx)": {"brs": len(brs6), "test_cases": 0},
        },
        "br_types": {},
        "output_files": {
            "brs": str(br_path),
            "test_cases": str(tc_path),
        }
    }
    
    # Count BR types
    for br in all_brs:
        t = br.get("type", "Unknown")
        summary["br_types"][t] = summary["br_types"].get(t, 0) + 1
    
    summary_path = OUTPUT_DIR / "summary.json"
    with open(summary_path, "w") as f:
        json.dump(summary, f, indent=2, ensure_ascii=False)
    
    # ── Print Summary ──
    print(f"\n{'=' * 70}")
    print(f"✅ SEPARATION COMPLETE")
    print(f"{'=' * 70}")
    print(f"\n📊 RESULTS:")
    print(f"  Total BRs extracted:        {len(all_brs)}")
    print(f"  Total Test Cases extracted:  {len(all_test_cases)}")
    print(f"    └─ With detailed steps:    {summary['test_cases_with_detailed_steps']}")
    print(f"\n📁 OUTPUT FILES:")
    print(f"  BRs:        {br_path}")
    print(f"  Test Cases: {tc_path}")
    print(f"  Summary:    {summary_path}")
    
    print(f"\n📋 BY SOURCE:")
    for src, counts in summary["by_source"].items():
        print(f"  {src}")
        print(f"    BRs: {counts['brs']}  |  Test Cases: {counts['test_cases']}")
    
    print(f"\n📋 BR TYPES:")
    for t, count in sorted(summary["br_types"].items()):
        print(f"  {t}: {count}")


if __name__ == "__main__":
    main()
