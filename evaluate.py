import json
import openpyxl
import os
import dspy
from pathlib import Path
from typing import List, Dict
from dotenv import load_dotenv

# Load environment
load_dotenv()

class RequirementMatcher(dspy.Signature):
    """Determine if a generated requirement matches a ground truth requirement.
    
    Match if they describe the SAME BUSINESS FUNCTION, even if wording differs.
    Ignore technical implementation details (e.g., if one mentions 'API' and the other doesn't).
    """
    extracted_title = dspy.InputField()
    extracted_desc = dspy.InputField()
    gt_title = dspy.InputField()
    gt_desc = dspy.InputField()
    
    matches = dspy.OutputField(desc="yes/no")
    reason = dspy.OutputField(desc="brief explanation")

def _get_lm():
    groq_key = os.getenv("GROQ_API_KEY")
    if not groq_key:
        raise ValueError("GROQ_API_KEY not found in .env")
    return dspy.LM('groq/llama-3.3-70b-versatile', api_key=groq_key)

def load_ground_truth(excel_path: str, sheet_name: str = "Requirements") -> List[Dict]:
    """Load ground truth from Excel."""
    wb = openpyxl.load_workbook(excel_path)
    ws = wb[sheet_name]
    
    requirements = []
    for row in range(3, ws.max_row + 1):
        req_id = ws.cell(row, 1).value
        req_name = ws.cell(row, 2).value
        req_desc = ws.cell(row, 3).value
        
        if req_id and req_name:
            requirements.append({
                'id': str(req_id),
                'name': str(req_name),
                'description': str(req_desc) if req_desc else ""
            })
    
    return requirements

def _tokenize(text: str) -> set:
    """Tokenize text into words, stripping punctuation and normalizing plurals."""
    import re
    tokens = set(re.findall(r'[a-z0-9](?:[a-z0-9_-]*[a-z0-9])?', text.lower()))
    
    # Plural normalization: strip trailing 's' (not 'ss', not short words)
    normalized = set()
    for t in tokens:
        if t.endswith('s') and not t.endswith('ss') and len(t) > 3:
            normalized.add(t[:-1])
        else:
            normalized.add(t)
    return normalized


def calculate_similarity(text1: str, text2: str) -> float:
    """Word overlap similarity (Jaccard) with punctuation handling."""
    words1 = _tokenize(text1)
    words2 = _tokenize(text2)

    if not words1 or not words2:
        return 0.0

    intersection = words1.intersection(words2)
    union = words1.union(words2)

    return len(intersection) / len(union)


def _title_similarity(title1: str, title2: str) -> float:
    """Compare two titles using multiple strategies."""
    t1 = title1.lower().strip()
    t2 = title2.lower().strip()

    # Exact match
    if t1 == t2:
        return 1.0

    # One title contains the other (substring)
    if t1 in t2 or t2 in t1:
        return 0.92

    # Word containment: all words of shorter title appear in longer title
    words1 = _tokenize(t1)
    words2 = _tokenize(t2)
    shorter, longer = (words1, words2) if len(words1) <= len(words2) else (words2, words1)
    if shorter and shorter.issubset(longer):
        return 0.88

    # Jaccard on title words
    if not words1 or not words2:
        return 0.0
    intersection = words1.intersection(words2)
    union = words1.union(words2)
    return len(intersection) / len(union)

def evaluate(extracted: List[Dict], ground_truth: List[Dict], threshold: float = 0.3) -> Dict:
    """Evaluate extracted vs ground truth using exclusive matching to ensure accurate metrics."""
    
    unique_matched_gt = set()      # Unique GT IDs that were found
    correct_extractions = []       # List of extractions that successfully found a NEW GT
    redundant_extractions = []     # List of extractions that matched an already-found GT
    over_created_extractions = []  # List of extractions with no match at all
    
    lm = _get_lm()
    matcher = dspy.Predict(RequirementMatcher)
    
    print(f"  Analysing {len(extracted)} extracted requirements against {len(ground_truth)} ground truth items...")

    # Two-pass matching: prefer unmatched GTs to maximize recall
    import re as re_mod

    def _enrich_desc(ext):
        """Build enriched description including acceptance criteria."""
        desc = ext.get('description', '')
        ac = ext.get('acceptance_criteria', [])
        if ac:
            ac_text = ". ".join(str(a) for a in ac)
            desc = f"{desc}. Acceptance criteria: {ac_text}"
        return desc

    def _find_match(ext, gt_candidates, lm_matcher):
        """Find the best matching GT from candidates."""
        ext_desc = _enrich_desc(ext)
        scored = []
        for gt in gt_candidates:
            title_sim = _title_similarity(ext['title'], gt['name'])
            desc_sim = calculate_similarity(ext_desc, gt['description'])
            best_score = max(title_sim, desc_sim)
            scored.append((gt, title_sim, desc_sim, best_score))
        scored.sort(key=lambda x: x[3], reverse=True)

        for gt, title_sim, desc_sim, best_score in scored:
            # Strategy 1: Title match (high similarity or substring)
            t1, t2 = ext['title'].lower(), gt['name'].lower()
            if title_sim >= 0.65 or t1 in t2 or t2 in t1:
                return gt

            # Strategy 2: Shared Domain Keywords (Very Strong Signal)
            high_value_keywords = {'p2p', 'cashback', 'crypto', 'biometric', 'onfido', 'plaid', 'kanban', 'dedup', 'sla', 'hy-yield', 'fractional', 'trading', 'banking'}
            ext_text = f"{ext['title']} {ext_desc}".lower()
            gt_text = f"{gt['name']} {gt['description']}".lower()
            
            shared_domain = _tokenize(ext_text) & _tokenize(gt_text) & high_value_keywords
            if shared_domain and (title_sim >= 0.20 or desc_sim >= 0.10):
                return gt


            # Strategy 3: Moderate overlap
            if title_sim >= 0.45 and desc_sim >= 0.20:
                return gt

            # Strategy 4: Shared domain identifiers
            ext_ids = set(re_mod.findall(r'[A-Z][A-Z0-9_]+-\d+', f"{ext['title']} {ext_desc}"))
            gt_ids = set(re_mod.findall(r'[A-Z][A-Z0-9_]+-\d+', f"{gt['name']} {gt['description']}"))
            if ext_ids and gt_ids and ext_ids & gt_ids:
                return gt

            # Skip if no overlap
            if best_score < 0.04: # Lowered from 0.08 to allow more LLM checks
                continue

            # Strategy 5: LLM check
            result = lm_matcher(
                extracted_title=ext['title'],
                extracted_desc=ext_desc,
                gt_title=gt['name'],
                gt_desc=gt['description']
            )

            if result.matches.lower().strip() == 'yes':
                return gt
        return None

    with dspy.context(lm=lm):
        # Pass 1: Match each extracted item to UNMATCHED GTs only (maximize TP)
        ext_matches = {}  # ext_index -> matched GT
        for i, ext in enumerate(extracted):
            print(f"    [{i+1}/{len(extracted)}] Matching: {ext['title'][:40]}...", end='\r')
            unmatched_gts = [gt for gt in ground_truth if gt['id'] not in unique_matched_gt]
            match = _find_match(ext, unmatched_gts, matcher)
            if match:
                ext_matches[i] = match
                unique_matched_gt.add(match['id'])
                correct_extractions.append({
                    'extracted_id': ext['requirement_id'],
                    'extracted_title': ext['title'],
                    'gt_id': match['id'],
                    'gt_name': match['name'],
                    'semantic_match': True
                })

        # Pass 2: For unmatched extracted items, check if they match already-matched GTs (redundant)
        for i, ext in enumerate(extracted):
            if i in ext_matches:
                continue  # Already matched in pass 1
            matched_gts = [gt for gt in ground_truth if gt['id'] in unique_matched_gt]
            match = _find_match(ext, matched_gts, matcher)
            if match:
                redundant_extractions.append({
                    'extracted_id': ext['requirement_id'],
                    'extracted_title': ext['title'],
                    'gt_id': match['id'],
                    'gt_name': match['name'],
                    'semantic_match': True
                })
            else:
                over_created_extractions.append({
                    'extracted_id': ext['requirement_id'],
                    'title': ext['title'],
                    'type': ext.get('type', 'Unknown')
                })
    
    print("\n  Matching complete.")
    
    # Step 3: Identify missed requirements (False Negatives)
    missed = []
    for gt in ground_truth:
        if gt['id'] not in unique_matched_gt:
            missed.append({
                'gt_id': gt['id'],
                'gt_name': gt['name'],
                'description': gt['description']
            })
    
    # Step 4: Calculate Standard Metrics
    # TP = Unique GTs found
    # FP = Junk extractions + Redundant extractions
    # FN = GTs missed
    tp = len(unique_matched_gt)
    fp = len(over_created_extractions) + len(redundant_extractions)
    fn = len(missed)
    
    precision = tp / (tp + fp) if (tp + fp) > 0 else 0
    recall = tp / (tp + fn) if (tp + fn) > 0 else 0
    f1 = 2 * precision * recall / (precision + recall) if (precision + recall) > 0 else 0
    
    return {
        'metrics': {
            'precision': precision,
            'recall': recall,
            'f1_score': f1
        },
        'counts': {
            'gt': len(ground_truth),
            'extracted': len(extracted),
            'true_positives': tp,
            'false_negatives': fn,
            'false_positives': fp,
            'redundant': len(redundant_extractions)
        },
        'matched': correct_extractions,
        'missed': missed,
        'over_created': over_created_extractions,
        'redundant_items': redundant_extractions
    }

def main():
    print("=" * 80)
    print("EVALUATION")
    print("=" * 80)

    # Load config
    config_file = Path("config/config.json")
    with open(config_file) as f:
        config = json.load(f)

    # Read current project from multi-project config
    current_project = config.get('current_project')
    if not current_project:
        print("ERROR: 'current_project' not set in config.json")
        return

    projects = config.get('projects', {})
    if current_project not in projects:
        print(f"ERROR: Project '{current_project}' not found in config.projects")
        return

    project_config = projects[current_project]
    output_dir = Path(config.get('output_dir', 'output'))

    # Check ground truth
    gt_path = project_config.get('ground_truth')
    if not gt_path:
        print(f"No ground truth file configured for project '{current_project}'.")
        print("Add a 'ground_truth' path in config.json to enable evaluation.")
        return

    gt_file = Path(gt_path)
    req_file = output_dir / f"{current_project}_requirements.json"

    if not gt_file.exists():
        print(f"ERROR: Ground truth not found: {gt_file}")
        return

    if not req_file.exists():
        print(f"ERROR: Requirements not found: {req_file}")
        print("Run extract_requirements.py first.")
        return

    # Load
    print(f"\nProject: {current_project}")
    print(f"\nLoading ground truth: {gt_file}")
    ground_truth = load_ground_truth(str(gt_file))
    print(f"Loaded {len(ground_truth)} requirements")

    print(f"\nLoading extracted requirements: {req_file}")
    with open(req_file) as f:
        data = json.load(f)
        extracted = data['requirements']
    print(f"Loaded {len(extracted)} requirements")

    # Evaluate
    print(f"\nEvaluating...")
    evaluation = evaluate(extracted, ground_truth)

    # Print results
    print(f"\n{'=' * 80}")
    print("RESULTS:")
    print(f"{'=' * 80}")
    print(f"Precision:  {evaluation['metrics']['precision']:.2%}")
    print(f"Recall:     {evaluation['metrics']['recall']:.2%}")
    print(f"F1 Score:   {evaluation['metrics']['f1_score']:.2%}")
    print(f"")
    print(f"Ground Truth:      {evaluation['counts']['gt']}")
    print(f"Extracted:         {evaluation['counts']['extracted']}")
    print(f"True Positives:    {evaluation['counts']['true_positives']}")
    print(f"False Negatives:   {evaluation['counts']['false_negatives']} (missed)")
    print(f"False Positives:   {evaluation['counts']['false_positives']} (over-extracted)")
    print(f"{'=' * 80}")

    # Save
    eval_file = output_dir / f"{current_project}_evaluation.json"
    with open(eval_file, 'w') as f:
        json.dump(evaluation, f, indent=2)

    print(f"\nSaved to: {eval_file}")
    print("Done!")

if __name__ == "__main__":
    main()
