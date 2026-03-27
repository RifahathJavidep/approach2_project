import base64
from typing import List, Dict, Any

from docx import Document


class DOCXExtractor:
    """Extract text and images from Word documents. EMF/WMF images are skipped."""

    def extract_text(self, file_path: str) -> str:
        try:
            doc = Document(file_path)
            text = "\n".join([para.text for para in doc.paragraphs if para.text.strip()])

            if text:
                print(f"  ✓ DOCX text extracted ({len(text)} chars, {len(doc.paragraphs)} paragraphs)")
            else:
                print(f"  ⚠ DOCX has no paragraph text")

            return text

        except Exception as e:
            print(f"  ✗ Error extracting text from DOCX: {e}")
            return ""

    def extract_images(self, file_path: str) -> List[Dict[str, Any]]:
        images = []
        try:
            doc = Document(file_path)
            img_counter = 0

            for rel in doc.part.rels.values():
                if "image" not in rel.reltype:
                    continue

                try:
                    image_part = rel.target_part
                    image_bytes = image_part.blob

                    if len(image_bytes) < 1000:
                        continue

                    content_type = image_part.content_type
                    image_ext = content_type.split("/")[-1]

                    format_map = {
                        "jpg": "jpeg",
                        "jpe": "jpeg",
                        "x-png": "png",
                    }
                    img_format = format_map.get(image_ext.lower(), image_ext.lower())

                    if img_format in {"x-emf", "x-wmf", "emf", "wmf"}:
                        continue

                    img_counter += 1

                    images.append({
                        "rel_id": rel.rId,
                        "index": img_counter,
                        "format": img_format,
                        "base64": base64.b64encode(image_bytes).decode("utf-8"),
                        "bytes": image_bytes,
                        "source": f"docx_img{img_counter}.{img_format}",
                        "size_bytes": len(image_bytes),
                    })

                except Exception as e:
                    print(f"   ⚠ Could not extract image {rel.rId}: {e}")
                    continue

        except Exception as e:
            print(f"   ✗ Error extracting images from DOCX: {e}")

        if images:
            print(f"  ✓ Extracted {len(images)} embedded images from DOCX")

        return images
