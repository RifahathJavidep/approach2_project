"""
Test Generation Pipeline — Phase 2 Orchestrator

Takes requirements from Phase 1 and produces:
  1. Test plan JSON (structured test cases)
  2. Test plan Excel (formatted for UAT sign-off)

Usage:
    from testgen import TestGenPipeline

    pipeline = TestGenPipeline()

    # Requirements only
    result = pipeline.run(requirements, project_name="ptw")

    # Requirements + Document context (richer test cases)
    result = pipeline.run(requirements, project_name="ptw", document_paths=["doc.pdf"])
"""

import json
import os
from pathlib import Path
from typing import List, Dict, Any, Optional

from .planner import TestCasePlanner
from .exporters import export_test_plan_to_excel


def _extract_text_from_file(file_path: str) -> str:
    """Extract text from a document file for context injection."""
    ext = Path(file_path).suffix.lower()

    if ext == ".pdf":
        try:
            import pymupdf4llm
            text = pymupdf4llm.to_markdown(file_path)
            if text and len(text.strip()) > 100:
                return text
        except Exception:
            pass

        # Fallback to PyMuPDF
        try:
            import fitz
            doc = fitz.open(file_path)
            text = "\n".join([page.get_text() for page in doc])
            doc.close()
            return text
        except Exception:
            return ""

    elif ext == ".docx":
        try:
            from docx import Document
            doc = Document(file_path)
            return "\n".join([p.text for p in doc.paragraphs if p.text.strip()])
        except Exception:
            return ""

    elif ext == ".pptx":
        try:
            from pptx import Presentation
            prs = Presentation(file_path)
            texts = []
            for slide in prs.slides:
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text.strip():
                        texts.append(shape.text)
            return "\n".join(texts)
        except Exception:
            return ""

    elif ext == ".txt" or ext == ".md":
        with open(file_path, "r", encoding="utf-8") as f:
            return f.read()

    return ""


class TestGenPipeline:
    """
    Phase 2 Orchestrator.

    Input:  List of requirement dicts (from Phase 1) + optional document paths
    Output: Test plan JSON + Excel file
    """

    def __init__(self, model: str = "llama-3.3-70b-versatile"):
        self.planner = TestCasePlanner(model=model)

    def run(
        self,
        requirements: List[Dict[str, Any]],
        project_name: str = "project",
        output_dir: str = "output",
        document_paths: Optional[List[str]] = None,
        source_texts: Optional[Dict[str, str]] = None,
        status_callback: Optional[callable] = None,
    ) -> Dict[str, Any]:
        """
        Run the full test generation pipeline.

        Args:
            requirements: List of requirement dicts from Phase 1
            project_name: Project name (used for file naming)
            output_dir: Directory to save output files
            document_paths: Optional list of document file paths to extract
                           text from for richer test case generation
            source_texts: Optional pre-extracted text dict (filename -> text).
                         If provided, document_paths is ignored.
            status_callback: Optional fn(message) for progress updates

        Returns:
            Dict with test_plan, json_path, excel_path
        """
        os.makedirs(output_dir, exist_ok=True)

        # ── Step 0: Load document context ───────────────────────────────
        if source_texts is None and document_paths:
            source_texts = {}
            print(f"\n  Loading document context for richer test cases...")
            for doc_path in document_paths:
                if os.path.exists(doc_path):
                    filename = Path(doc_path).name
                    text = _extract_text_from_file(doc_path)
                    if text:
                        source_texts[filename] = text
                        print(f"    ✓ {filename}: {len(text):,} chars loaded")
                    else:
                        print(f"    ⚠ {filename}: No text extracted")
                else:
                    print(f"    ✗ {doc_path}: File not found")

        # ── Step 1: Generate test cases ─────────────────────────────────
        test_plan = self.planner.generate(
            requirements=requirements,
            project_name=project_name,
            source_texts=source_texts,
            status_callback=status_callback,
        )

        # ── Step 2: Save JSON ───────────────────────────────────────────
        json_path = os.path.join(output_dir, f"{project_name}_test_plan.json")
        with open(json_path, "w", encoding="utf-8") as f:
            json.dump(test_plan, f, indent=2, ensure_ascii=False)
        print(f"  ✓ JSON saved: {json_path}")

        # ── Step 3: Export to Excel ─────────────────────────────────────
        excel_path = os.path.join(output_dir, f"{project_name}_test_plan.xlsx")
        export_test_plan_to_excel(test_plan, excel_path, project_name)

        return {
            "test_plan": test_plan,
            "json_path": json_path,
            "excel_path": excel_path,
            "total_requirements": test_plan["total_requirements"],
            "total_test_cases": test_plan["total_test_cases"],
        }

    def run_from_json(
        self,
        json_path: str,
        project_name: str = None,
        output_dir: str = "output",
        document_paths: Optional[List[str]] = None,
    ) -> Dict[str, Any]:
        """
        Convenience method: load requirements from a JSON file and run.

        Args:
            json_path: Path to requirements JSON file
            project_name: Project name (auto-detected from filename if None)
            output_dir: Directory to save output files
            document_paths: Optional list of source document paths for context

        Returns:
            Dict with test_plan, json_path, excel_path
        """
        with open(json_path, "r", encoding="utf-8") as f:
            requirements = json.load(f)

        if not isinstance(requirements, list):
            raise ValueError(f"Expected a JSON array of requirements, got {type(requirements).__name__}")

        if project_name is None:
            # Auto-detect from filename: "ptw_requirements.json" → "ptw"
            name = Path(json_path).stem
            project_name = name.replace("_requirements", "").replace("_reqs", "")

        return self.run(
            requirements, project_name, output_dir,
            document_paths=document_paths,
        )
