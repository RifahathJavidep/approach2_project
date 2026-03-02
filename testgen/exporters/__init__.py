"""
Excel Exporter for Test Plans

Exports AI-generated test plans to a professionally formatted Excel file
suitable for UAT sign-off and test management.

Columns:
    A: Requirement ID
    B: Requirement Title
    C: Test Case ID
    D: Test Case Title
    E: Test Case Description
    F: Test Type
    G: Test Phase
    H: Priority
    I: Prerequisites
    J: Test Steps
    K: Expected Result
    L: Test Data
    M: Status
"""

import os
from typing import Dict, Any

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

from ..utils import to_excel_safe


def export_test_plan_to_excel(
    test_plan: Dict[str, Any],
    output_path: str,
    project_name: str = "Project",
) -> str:
    """
    Export test plan to a professionally formatted Excel file.

    Args:
        test_plan: Test plan dict from TestCasePlanner.generate()
        output_path: Path to save the Excel file
        project_name: Project name for the header

    Returns:
        Path to the saved Excel file
    """
    print(f"\n  Exporting test plan to Excel...")

    wb = Workbook()
    ws = wb.active
    ws.title = "Test Cases"

    # ── Styles ──────────────────────────────────────────────────────────
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=11, name="Calibri")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

    cell_font = Font(size=10, name="Calibri")
    cell_align = Alignment(vertical="top", wrap_text=True)
    cell_align_center = Alignment(horizontal="center", vertical="top", wrap_text=True)

    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )

    # Priority color fills
    priority_fills = {
        "Critical": PatternFill(start_color="FFE0E0", fill_type="solid"),
        "High": PatternFill(start_color="FFF3E0", fill_type="solid"),
        "Medium": PatternFill(start_color="E8F5E9", fill_type="solid"),
        "Low": PatternFill(start_color="E3F2FD", fill_type="solid"),
    }

    # ── Headers ─────────────────────────────────────────────────────────
    headers = [
        "Requirement ID",
        "Requirement Title",
        "Test Case ID",
        "Test Case Title",
        "Test Case Description",
        "Test Type",
        "Test Phase",
        "Priority",
        "Prerequisites",
        "Test Steps",
        "Expected Result",
        "Test Data",
        "Status",
    ]

    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num)
        cell.value = header
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_align
        cell.border = thin_border

    ws.freeze_panes = "A2"

    # ── Data Rows ───────────────────────────────────────────────────────
    row_num = 2
    global_tc_counter = 1

    for plan in test_plan.get("test_plans", []):
        req_id = plan.get("requirement_id", "N/A")
        req_title = plan.get("requirement_title", "Untitled")

        for tc in plan.get("test_cases", []):
            tc_id = f"TC-{global_tc_counter:03d}"
            global_tc_counter += 1

            # Format test steps as numbered list
            steps = tc.get("test_steps", [])
            if isinstance(steps, list):
                steps_text = "\n".join([f"{i+1}. {s}" for i, s in enumerate(steps)])
            else:
                steps_text = str(steps)

            row_data = [
                req_id,
                to_excel_safe(req_title),
                tc_id,
                to_excel_safe(tc.get("title", "Untitled")),
                to_excel_safe(tc.get("description", "")),
                to_excel_safe(tc.get("test_type", "Functional")),
                to_excel_safe(tc.get("test_phase", "E2E")),
                to_excel_safe(tc.get("priority", "Medium")),
                to_excel_safe(tc.get("prerequisites", "")),
                steps_text,
                to_excel_safe(tc.get("expected_result", "")),
                to_excel_safe(tc.get("test_data", "")),
                to_excel_safe(tc.get("status", "Not Executed")),
            ]

            for col_num, value in enumerate(row_data, 1):
                cell = ws.cell(row=row_num, column=col_num)
                cell.value = value
                cell.font = cell_font
                cell.border = thin_border

                # Center-align short columns
                if col_num in [1, 3, 6, 7, 8, 13]:
                    cell.alignment = cell_align_center
                else:
                    cell.alignment = cell_align

            # Apply priority color
            priority = tc.get("priority", "Medium")
            priority_cell = ws.cell(row=row_num, column=8)
            if priority in priority_fills:
                priority_cell.fill = priority_fills[priority]

            # Alternate row shading
            if row_num % 2 == 0:
                alt_fill = PatternFill(start_color="F8F9FA", fill_type="solid")
                for col in range(1, len(headers) + 1):
                    cell = ws.cell(row=row_num, column=col)
                    if col != 8:  # Don't override priority color
                        cell.fill = alt_fill

            row_num += 1

    # ── Column Widths ───────────────────────────────────────────────────
    widths = {
        1: 18,   # Requirement ID
        2: 35,   # Requirement Title
        3: 12,   # Test Case ID
        4: 35,   # TC Title
        5: 55,   # TC Description
        6: 14,   # Test Type
        7: 12,   # Test Phase
        8: 10,   # Priority
        9: 40,   # Prerequisites
        10: 60,  # Test Steps
        11: 45,  # Expected Result
        12: 30,  # Test Data
        13: 14,  # Status
    }

    for col_num, width in widths.items():
        ws.column_dimensions[get_column_letter(col_num)].width = width

    ws.row_dimensions[1].height = 35

    # ── Save ────────────────────────────────────────────────────────────
    os.makedirs(os.path.dirname(output_path) if os.path.dirname(output_path) else ".", exist_ok=True)
    wb.save(output_path)

    total_tests = row_num - 2
    print(f"  ✓ Excel saved: {output_path}")
    print(f"    Total test cases: {total_tests}")
    print(f"    Requirements covered: {len(test_plan.get('test_plans', []))}")

    return output_path
