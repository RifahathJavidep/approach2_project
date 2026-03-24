"""
Excel Exporter for Test Plans

Exports AI-generated test plans to a professionally formatted Excel file
suitable for UAT sign-off and test management.

Columns:
    A: Requirement ID
    B: Requirement Title
    C: Test Case Title
    D: Test Case Description
    E: Test Phases
    F: Test Types
    G: Prerequisites
    H: Test Steps
    I: Expected Result
    J: Status
"""

import os
from typing import Dict, Any

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

from ..utils import to_excel_safe

def export_test_plan_to_excel(
    test_plan: Dict[str, Any],
    output_path: str,
    project_name: str = "Project",
) -> str:
    """
    Export test plan to a professionally formatted Excel file.

    Args:
        test_plan: Test plan dict from TestCasePlanner.generate()
        output_path: Path to save the Excel file
        project_name: Project name for the header

    Returns:
        Path to the saved Excel file
    """
    print(f"\n  Exporting test plan to Excel...")

    wb = Workbook()
    ws = wb.active
    ws.title = "Test Cases"

    # ── Styles ──────────────────────────────────────────────────────────
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=11, name="Calibri")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

    cell_font = Font(size=10, name="Calibri")
    cell_align = Alignment(vertical="top", wrap_text=True)
    cell_align_center = Alignment(horizontal="center", vertical="top", wrap_text=True)

    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )

    # ── Headers ─────────────────────────────────────────────────────────
    headers = [
        "Requirement ID",
        "Requirement Title",
        "Test Case Title",
        "Test Case Description",
        "Test Phases",
        "Test Types",
        "Prerequisites",
        "Test Steps",
        "Expected Result",
        "Status",
    ]

    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num)
        cell.value = header
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_align
        cell.border = thin_border

    ws.freeze_panes = "A2"

    # ── Data Rows ───────────────────────────────────────────────────────
    row_num = 2
    global_tc_counter = 1

    for plan in test_plan.get("test_plans", []):
        req_id = plan.get("requirement_id", "N/A")
        req_title = plan.get("requirement_title", "Untitled")

        for tc in plan.get("test_cases", []):
            global_tc_counter += 1

            # Format testCaseSteps as numbered list
            tc_steps = tc.get("testCaseSteps", [])
            if isinstance(tc_steps, list):
                steps_lines = []
                for s in tc_steps:
                    order = s.get("orderNumber", 0) + 1
                    desc = s.get("description", "")
                    exp = s.get("expectedResult", "")
                    line = f"{order}. {desc}"
                    if exp:
                        line += f"\n   → {exp}"
                    steps_lines.append(line)
                steps_text = "\n".join(steps_lines)
            else:
                steps_text = str(tc_steps)

            # Format enum arrays as comma-separated strings
            test_phases = ", ".join(tc.get("testPhases", []))
            test_types = ", ".join(tc.get("testTypes", []))

            row_data = [
                to_excel_safe(tc.get("requirementId", req_id)),
                to_excel_safe(req_title),
                to_excel_safe(tc.get("title", "Untitled")),
                to_excel_safe(tc.get("description", "")),
                to_excel_safe(test_phases),
                to_excel_safe(test_types),
                to_excel_safe(tc.get("prerequisites", "")),
                steps_text,
                to_excel_safe(tc.get("expectedResult", "")),
                to_excel_safe(tc.get("status", "DRAFT")),
            ]

            for col_num, value in enumerate(row_data, 1):
                cell = ws.cell(row=row_num, column=col_num)
                cell.value = value
                cell.font = cell_font
                cell.border = thin_border

                # Center-align short columns
                if col_num in [1, 5, 6, 10]:
                    cell.alignment = cell_align_center
                else:
                    cell.alignment = cell_align

            # Alternate row shading
            if row_num % 2 == 0:
                alt_fill = PatternFill(start_color="F8F9FA", fill_type="solid")
                for col in range(1, len(headers) + 1):
                    ws.cell(row=row_num, column=col).fill = alt_fill

            row_num += 1

    # ── Column Widths ───────────────────────────────────────────────────
    widths = {
        1: 18,   # Requirement ID
        2: 35,   # Requirement Title
        3: 40,   # TC Title
        4: 55,   # TC Description
        5: 28,   # Test Phases
        6: 18,   # Test Types
        7: 45,   # Prerequisites
        8: 65,   # Test Steps
        9: 45,   # Expected Result
        10: 14,  # Status
    }

    for col_num, width in widths.items():
        ws.column_dimensions[get_column_letter(col_num)].width = width

    ws.row_dimensions[1].height = 35

    # ── Save ────────────────────────────────────────────────────────────
    os.makedirs(os.path.dirname(output_path) if os.path.dirname(output_path) else ".", exist_ok=True)
    wb.save(output_path)

    total_tests = row_num - 2
    print(f"  ✓ Excel saved: {output_path}")
    print(f"    Total test cases: {total_tests}")
    print(f"    Requirements covered: {len(test_plan.get('test_plans', []))}")

    return output_path
